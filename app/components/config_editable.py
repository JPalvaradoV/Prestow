"""
Lectura y escritura de las tablas editables (Buque, Productos, Viaje,
Rotación) que alimentan los formularios de la Fase 2.

Funciones puras: no acceden a st.session_state. Las páginas de Streamlit
las llaman y guardan el resultado en session_state.

Formato de intercambio: una carpeta con cuatro CSV (buque_*.csv,
productos_*.csv, rotacion_*.csv, viaje_*.csv), el mismo que ya consumen
packer_2d.leer_entrada y modelo_prestow.cargar_desde_csv. Así los cambios
del usuario se recalculan con el mismo código que el caso base, sin tocar
modelo_prestow.py ni packer_2d.py.
"""

from __future__ import annotations

import csv
from pathlib import Path

import pandas as pd

COLUMNAS_BUQUE = ["bodega", "largo_m", "ancho_m", "planes", "cuadrilla"]
COLUMNAS_PRODUCTOS = ["producto", "huella_largo_m", "huella_ancho_m", "peso_t", "origen_dato"]
COLUMNAS_ROTACION = ["destino", "orden_descarga", "nota"]
COLUMNAS_VIAJE = ["producto", "destino", "unidades"]

# Lista de referencia de puertos que reciben celulosa/madera en rollo, para el
# desplegable de Rotación. NO es un dato verificado por el puerto ni por el
# modelo: es una lista de conveniencia compilada por región para evitar
# errores de tipeo al escribir el nombre. Incluye los cuatro puertos del caso
# base (TAICHUNG, QINGDAO, KUNSAN, ULSAN). Formato: nombre de la ciudad
# portuaria en mayúsculas, sin espacios, igual que los datos existentes. Si
# falta un puerto, se agrega a esta lista.
PUERTOS_CELULOSA_REFERENCIA = [
    # Taiwán
    "TAICHUNG", "KAOHSIUNG",
    # China
    "QINGDAO", "SHANGHAI", "NINGBO", "ZHANGJIAGANG", "NANTONG",
    "LIANYUNGANG", "XINGANG", "DALIAN",
    # Corea del Sur
    "ULSAN", "KUNSAN", "BUSAN", "INCHEON",
    # Japón
    "YOKOHAMA", "NAGOYA", "OSAKA",
    # India
    "CHENNAI", "MUNDRA",
    # Europa
    "ROTTERDAM", "ANTWERP",
]


def _filas_hoja(ws, columnas: list[str]) -> list[dict]:
    """Lee una hoja de Excel saltando títulos, hasta la primera fila vacía."""
    fila_enc = None
    for r in range(1, 12):
        if str(ws.cell(row=r, column=1).value or "").strip() == columnas[0]:
            fila_enc = r
            break
    if fila_enc is None:
        raise ValueError(f"En la hoja '{ws.title}' no se encontró la columna '{columnas[0]}'")

    salida = []
    r = fila_enc + 1
    while True:
        primera = ws.cell(row=r, column=1).value
        if primera is None or str(primera).strip() == "" or str(primera).strip().upper() == "TOTAL":
            break
        salida.append({col: ws.cell(row=r, column=j).value for j, col in enumerate(columnas, start=1)})
        r += 1
    return salida


def leer_tablas_editables(ruta_datos: str | Path) -> dict[str, pd.DataFrame]:
    """
    Lee las cuatro tablas (Buque, Productos, Rotación, Viaje) del Excel de
    entrada como DataFrames listos para editar en st.data_editor.
    """
    from openpyxl import load_workbook

    wb = load_workbook(ruta_datos, data_only=True)

    faltantes = [h for h in ("Buque", "Productos", "Rotacion", "Viaje") if h not in wb.sheetnames]
    if faltantes:
        raise ValueError(f"Al archivo le faltan estas hojas: {', '.join(faltantes)}")

    buque = pd.DataFrame(_filas_hoja(wb["Buque"], COLUMNAS_BUQUE))
    productos = pd.DataFrame(_filas_hoja(wb["Productos"], COLUMNAS_PRODUCTOS))
    rotacion = pd.DataFrame(_filas_hoja(wb["Rotacion"], COLUMNAS_ROTACION))
    viaje = pd.DataFrame(_filas_hoja(wb["Viaje"], COLUMNAS_VIAJE))

    buque["bodega"] = buque["bodega"].astype(int)
    buque["planes"] = buque["planes"].astype(int)
    buque["cuadrilla"] = buque["cuadrilla"].astype(int)
    # El Excel del caso base lista las bodegas de la 8 a la 1 (así las anota
    # el puerto); para editar es mas intuitivo de la 1 a la 8.
    buque = buque.sort_values("bodega").reset_index(drop=True)
    rotacion["orden_descarga"] = rotacion["orden_descarga"].astype(int)
    rotacion = rotacion.sort_values("orden_descarga").reset_index(drop=True)
    viaje["unidades"] = viaje["unidades"].astype(int)

    return {"buque": buque, "productos": productos, "rotacion": rotacion, "viaje": viaje}


def validar_tablas(tablas: dict[str, pd.DataFrame]) -> list[str]:
    """
    Validaciones mínimas antes de guardar: sin ellas el packer o el modelo
    fallan con errores poco claros. Devuelve la lista de problemas
    encontrados (vacía si todo está bien).
    """
    errores: list[str] = []
    buque, productos, rotacion, viaje = (
        tablas["buque"], tablas["productos"], tablas["rotacion"], tablas["viaje"]
    )

    if buque.empty:
        errores.append("El buque necesita al menos una bodega.")
    elif buque["bodega"].duplicated().any():
        errores.append("Hay números de bodega repetidos.")
    elif (buque["largo_m"] <= 0).any() or (buque["ancho_m"] <= 0).any():
        errores.append("Las dimensiones del piso deben ser mayores a 0.")
    elif (buque["planes"] <= 0).any():
        errores.append("El número de planes debe ser mayor a 0.")

    if productos.empty:
        errores.append("Debe haber al menos un producto.")
    elif productos["producto"].astype(str).str.strip().duplicated().any():
        errores.append("Hay productos con el mismo nombre.")
    elif (productos["huella_largo_m"] <= 0).any() or (productos["huella_ancho_m"] <= 0).any():
        errores.append("La huella de cada producto debe ser mayor a 0.")

    if rotacion.empty:
        errores.append("Debe haber al menos un destino en la rotación.")
    elif rotacion["destino"].astype(str).str.strip().str.upper().duplicated().any():
        errores.append("Hay destinos repetidos en la rotación.")
    elif rotacion["orden_descarga"].duplicated().any():
        errores.append("Dos destinos no pueden tener el mismo orden de descarga.")

    productos_validos = set(productos["producto"].astype(str).str.strip()) if not productos.empty else set()
    destinos_validos = set(rotacion["destino"].astype(str).str.strip().str.upper()) if not rotacion.empty else set()
    if viaje.empty:
        errores.append("El viaje necesita al menos una combinación producto-destino.")
    else:
        if (viaje["unidades"] < 0).any():
            errores.append("Las unidades del viaje no pueden ser negativas.")
        productos_viaje = set(viaje["producto"].astype(str).str.strip())
        destinos_viaje = set(viaje["destino"].astype(str).str.strip().str.upper())
        faltan_prod = productos_viaje - productos_validos
        faltan_dest = destinos_viaje - destinos_validos
        if faltan_prod:
            errores.append(f"El viaje usa productos que no están en Productos: {', '.join(sorted(faltan_prod))}")
        if faltan_dest:
            errores.append(f"El viaje usa destinos que no están en Rotación: {', '.join(sorted(faltan_dest))}")

    return errores


def guardar_tablas_editables(tablas: dict[str, pd.DataFrame], carpeta: str | Path) -> Path:
    """
    Escribe las cuatro tablas como CSV en `carpeta`, con el formato que
    esperan packer_2d.leer_entrada y modelo_prestow.cargar_desde_csv.
    Devuelve la carpeta.
    """
    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)

    def escribir(nombre: str, df: pd.DataFrame, columnas: list[str]) -> None:
        # to_dict("records") preserva el tipo de cada columna; iterrows() no
        # sirve aca porque sube todo a float64 cuando hay columnas mixtas
        # (ej. bodega=int junto a largo_m=float), y bodega=8 se escribe "8.0".
        ruta = carpeta / nombre
        with open(ruta, "w", newline="", encoding="utf-8") as f:
            wr = csv.DictWriter(f, fieldnames=columnas)
            wr.writeheader()
            for fila in df.to_dict("records"):
                wr.writerow({c: fila[c] for c in columnas})

    escribir("buque_editado.csv", tablas["buque"], COLUMNAS_BUQUE)
    escribir("productos_editado.csv", tablas["productos"], COLUMNAS_PRODUCTOS)
    escribir("rotacion_editado.csv", tablas["rotacion"], COLUMNAS_ROTACION)
    escribir("viaje_editado.csv", tablas["viaje"], COLUMNAS_VIAJE)
    return carpeta


def recalcular_capacidades(carpeta: str | Path) -> Path:
    """
    Corre packer_2d sobre la carpeta editada y guarda capacidades.csv ahí
    mismo. Se llama después de cualquier cambio en Buque o Productos, porque
    la capacidad depende de la geometría de ambos.
    """
    try:
        from . import packer_2d  # type: ignore[import]
    except ImportError:
        import packer_2d  # type: ignore[import]

    carpeta = Path(carpeta)
    bodegas, productos = packer_2d.leer_entrada(carpeta)
    filas = packer_2d.generar_tabla(bodegas, productos, mostrar=False)
    ruta_salida = carpeta / "capacidades.csv"
    packer_2d.guardar_tabla(filas, str(ruta_salida))
    return ruta_salida


def resumen_metadata(tablas: dict[str, pd.DataFrame], nombre_buque: str = "Buque editado") -> dict:
    """Arma el dict de metadata que usa la web para las cards de resumen."""
    buque, productos, rotacion, viaje = (
        tablas["buque"], tablas["productos"], tablas["rotacion"], tablas["viaje"]
    )
    destinos_ordenados = (
        rotacion.sort_values("orden_descarga")["destino"].astype(str).str.strip().str.upper().tolist()
    )

    peso_por_producto = {
        str(r["producto"]).strip(): float(r["peso_t"]) for _, r in productos.iterrows()
    }
    toneladas = sum(
        int(r["unidades"]) * peso_por_producto.get(str(r["producto"]).strip(), 0.0)
        for _, r in viaje.iterrows()
    )

    return {
        "buque": nombre_buque,
        "operador": "",
        "bodegas": int(len(buque)),
        "unidades": int(viaje["unidades"].sum()),
        "toneladas": round(toneladas),
        "destinos": destinos_ordenados,
        "productos": productos["producto"].astype(str).str.strip().tolist(),
        "makespan_manual_h": None,
        "es_caso_base": False,
    }
